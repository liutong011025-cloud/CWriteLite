import json
from pathlib import Path
import openpyxl

folder = Path(__file__).resolve().parents[1] / 'outputs' / 'process-verification'
expected = json.loads((folder / 'export-expected-count.json').read_text())['count']
workbook = openpyxl.load_workbook(folder / 'full-export-check.xlsx', read_only=True)
rows = workbook.worksheets[0].iter_rows()
header = [cell.value for cell in next(rows)]
text_index = header.index('beforeText')
count = formulas = literal = 0
uids = set()
last_bulk = False
for row in rows:
    count += 1
    uids.add(row[0].value)
    last_bulk |= str(row[0].value).endswith('-bulk-10025')
    formulas += sum(cell.data_type == 'f' for cell in row)
    literal += row[text_index].value == '=1+1' and row[text_index].data_type == 's'
assert count == expected and len(uids) == expected
assert last_bulk and formulas == 0 and literal == 10025
print(dict(export_rows=count, expected=expected, duplicate_ids=count-len(uids), formula_cells=formulas, literal_formula_strings=literal, sheets=workbook.sheetnames))
workbook.close()
